"""Módulo 2 — Backend de inventario: Exportar/Importar Excel + vencimientos.

Librería: openpyxl (estándar de facto para .xlsx en el ecosistema Python;
soportada por Excel/LibreOffice/Google Sheets). Se elige frente a CSV porque
es un único archivo binario que preserva tipos y formato, y frente a pandas
porque NO agrega una dependencia pesada (numpy/pandas) para lo que es, en
esencia, lectura/escritura de celdas.

Endpoints añadidos al router de productos:
  GET  /productos/exportar   -> descarga un .xlsx con el catálogo.
  POST /productos/importar   -> carga masiva desde .xlsx (valida y reporta).

Notas de ingeniería:
  * El import es TOLERANTE: reporta fila por fila los errores (no aborta todo).
  * Se procesa en una sola transacción; si hay errores de BD, rollback total.
  * Se soporta upsert por `codigo_barras` (clave natural del comercio).
"""
from tools_upgrade.patcher import apply

V = "app/routers/productos.py"

# 1) Imports necesarios (io, openpyxl, UploadFile/File/Response).
apply(
    V,
    (
        'from decimal import Decimal\n'
        "from typing import List\n"
        "from uuid import UUID\n"
        "\n"
        'from fastapi import APIRouter, Depends, HTTPException, Query, status\n'
        "from sqlalchemy.orm import Session",
        'import io\n'
        "from datetime import datetime\n"
        "from decimal import Decimal, InvalidOperation\n"
        "from typing import List, Optional\n"
        "from uuid import UUID\n"
        "\n"
        "from fastapi import (\n"
        "    APIRouter,\n"
        "    Depends,\n"
        "    File,\n"
        "    HTTPException,\n"
        "    Query,\n"
        "    UploadFile,\n"
        "    status,\n"
        ")\n"
        "from fastapi.responses import StreamingResponse\n"
        "from sqlalchemy.orm import Session",
    ),
)

# 2) Insertar los endpoints antes del endpoint /{producto_id} (GET detalle),
#    para que las rutas estáticas "/exportar" y "/importar" no colisionen con
#    el path param "/{producto_id}" (FastAPI resuelve por orden de declaración).
ANCLA = '@router.get("/{producto_id}", response_model=ProductoOut, summary="Obtener un producto")'
EXPORT_IMPORT = '''# ========================= EXPORTAR / IMPORTAR XLSX =========================
# Columnas del Excel de inventario (orden y encabezado estable, ES).
COLUMNAS_XLSX = [
    "nombre",
    "codigo_barras",
    "categoria",
    "talle",
    "color",
    "tipo_unidad",
    "precio",
    "costo",
    "porcentaje_ganancia",
    "stock",
    "stock_minimo",
    "fecha_vencimiento",
    "activo",
]


@router.get(
    "/exportar",
    summary="Exportar el catálogo de productos a Excel (.xlsx)",
)
def exportar_productos(
    usuario: UsuarioActual = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Genera y descarga un .xlsx con TODO el catálogo (activos e inactivos).

    Pensado para respaldo, edición masiva y carga entre instalaciones. El
    archivo resultante es re-importable tal cual (columnas idénticas).
    """
    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Productos"
        ws.append(COLUMNAS_XLSX)

        productos = db.query(Producto).order_by(Producto.nombre.asc()).all()
        for p in productos:
            ws.append(
                [
                    p.nombre,
                    p.codigo_barras or "",
                    p.categoria or "",
                    p.talle or "",
                    p.color or "",
                    p.tipo_unidad or "unidad",
                    float(p.precio or 0),
                    float(p.costo or 0),
                    float(p.porcentaje_ganancia or 0),
                    float(p.stock or 0),
                    float(p.stock_minimo or 0),
                    (
                        p.fecha_vencimiento.strftime("%Y-%m-%d")
                        if p.fecha_vencimiento
                        else ""
                    ),
                    1 if p.activo else 0,
                ]
            )

        # Anchos de columna legibles.
        for i, _ in enumerate(COLUMNAS_XLSX, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        headers = {
            "Content-Disposition": "attachment; filename=inventario.xlsx",
        }
        return StreamingResponse(
            buffer,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers=headers,
        )
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error al exportar productos: {str(exc)}",
        )


def _a_decimal(valor, default: str = "0") -> Decimal:
    """Convierte una celda de Excel a Decimal de forma tolerante."""
    if valor is None or valor == "":
        return Decimal(default)
    try:
        return Decimal(str(valor).replace(",", ".").strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f"valor numérico inválido: {valor!r}")


def _a_fecha(valor) -> Optional[datetime]:
    """Convierte una celda a datetime (acepta date, datetime o 'YYYY-MM-DD')."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor
    try:
        return datetime.strptime(str(valor).strip()[:10], "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"fecha inválida (use AAAA-MM-DD): {valor!r}")


@router.post(
    "/importar",
    summary="Importar/actualizar productos desde Excel (.xlsx) (solo admin)",
)
async def importar_productos(
    archivo: UploadFile = File(..., description="Archivo .xlsx de inventario"),
    usuario: UsuarioActual = Depends(requiere_rol("admin")),
    db: Session = Depends(get_db),
):
    """Carga masiva de catálogo desde .xlsx (upsert por código de barras).

    Estrategia (FIRME pero NUNCA destructiva):
      - Si la fila trae `codigo_barras` y YA existe un producto con ese código,
        se ACTUALIZA (upsert) — clave natural del comercio.
      - Si no existe (o no trae código), se CREA un producto nuevo.
      - El import es TOLERANTE: las filas con error se reportan y se omiten;
        las válidas se aplican. NO revierte todo por una fila mala.
      - Se aplica todo en una transacción por lote (commit al final).

    Devuelve un resumen: creados, actualizados, omitidos y detalle de errores.
    """
    nombre_archivo = (archivo.filename or "").lower()
    if not nombre_archivo.endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Formato no soportado. Subí un archivo .xlsx de Excel.",
        )

    try:
        from openpyxl import load_workbook

        contenido = await archivo.read()
        wb = load_workbook(io.BytesIO(contenido), data_only=True)
        ws = wb.active

        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El archivo está vacío.",
            )

        # Mapear encabezados -> índice de columna (case-insensitive, tolerante).
        encabezados = [
            str(c).strip().lower() if c is not None else "" for c in filas[0]
        ]
        idx = {col: i for i, col in enumerate(encabezados) if col}

        def celda(fila, col):
            i = idx.get(col)
            return fila[i] if (i is not None and i < len(fila)) else None

        creados = 0
        actualizados = 0
        omitidos = 0
        errores: list[dict] = []

        for num, fila in enumerate(filas[1:], start=2):  # fila 1 = encabezado
            if not any(c not in (None, "") for c in fila):
                continue  # fila totalmente vacía
            try:
                nombre = (celda(fila, "nombre") or "").strip() if isinstance(
                    celda(fila, "nombre"), str
                ) else celda(fila, "nombre")
                if not nombre:
                    raise ValueError("'nombre' es obligatorio")

                precio = _a_decimal(celda(fila, "precio"))
                if precio <= 0:
                    raise ValueError("'precio' debe ser mayor a cero")

                tipo = (str(celda(fila, "tipo_unidad") or "unidad")).strip().lower()
                if tipo not in Producto.TIPOS_UNIDAD:
                    tipo = "unidad"

                codigo = celda(fila, "codigo_barras")
                codigo = str(codigo).strip() if codigo not in (None, "") else None

                datos = {
                    "nombre": nombre,
                    "categoria": (celda(fila, "categoria") or None),
                    "talle": (celda(fila, "talle") or None),
                    "color": (celda(fila, "color") or None),
                    "tipo_unidad": tipo,
                    "precio": precio,
                    "costo": _a_decimal(celda(fila, "costo")),
                    "porcentaje_ganancia": _a_decimal(
                        celda(fila, "porcentaje_ganancia")
                    ),
                    "stock": _a_decimal(celda(fila, "stock")),
                    "stock_minimo": _a_decimal(celda(fila, "stock_minimo")),
                    "fecha_vencimiento": _a_fecha(celda(fila, "fecha_vencimiento")),
                    "activo": str(celda(fila, "activo") or "1").strip()
                    not in ("0", "false", "False", "no", "NO", ""),
                }

                existente = (
                    db.query(Producto)
                    .filter(Producto.codigo_barras == codigo)
                    .first()
                    if codigo
                    else None
                )

                if existente is not None:
                    for k, v in datos.items():
                        setattr(existente, k, v)
                    existente.actualizado_en = now_utc()
                    actualizados += 1
                else:
                    db.add(Producto(codigo_barras=codigo, **datos))
                    creados += 1
            except ValueError as ve:
                omitidos += 1
                errores.append({"fila": num, "error": str(ve)})
            except Exception as e:  # fila problemática: reportar y continuar
                omitidos += 1
                errores.append({"fila": num, "error": str(e)})

        db.commit()
        return {
            "creados": creados,
            "actualizados": actualizados,
            "omitidos": omitidos,
            "errores": errores[:50],  # tope para no devolver respuestas gigantes
            "mensaje": (
                f"Importación finalizada: {creados} creados, "
                f"{actualizados} actualizados, {omitidos} omitidos."
            ),
        }
    except HTTPException:
        db.rollback()
        raise
    except Exception as exc:  # pragma: no cover
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error al importar productos: {str(exc)}",
        )


'''
apply(V, (ANCLA, EXPORT_IMPORT + ANCLA))

print("[productos export/import] OK")
