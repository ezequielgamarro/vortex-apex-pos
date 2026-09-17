"""Router de EXPORTACIÓN TOTAL de datos del comercio (respaldo completo).

Objetivo de negocio:
    Permitir al ADMINISTRADOR descargar TODA la información de su negocio en
    un único libro Excel (.xlsx) con varias hojas, INCLUSO si la licencia está
    vencida. Es la garantía de "no rehén de datos": los datos del comercio son
    del comercio y debe poder recuperarlos en cualquier momento.

Diseño (por qué es seguro):
    - Es un endpoint de LECTURA (GET). El middleware de licencia
      (`app/middleware/licencia.py` -> `requiere_bloqueo`) SOLO bloquea la
      ESCRITURA (POST/PUT/PATCH/DELETE) sobre /ventas, /arqueo y /productos.
      Por lo tanto este GET queda PERMITIDO aun con la licencia vencida,
      sin necesidad de modificar el middleware.
    - La ruta NO vive bajo los prefijos bloqueados (se expone como
      `/exportar-todo`), reforzando que sea independiente del bloqueo.
    - Solo el rol 'admin' puede invocarla (`requiere_rol("admin")`): el cajero
      no exporta el respaldo completo (decisión de negocio del comercio).

Hoja por tabla (decisiones de contenido):
    - Productos : catálogo completo (activos e inactivos).
    - Ventas    : historial con artículos agregados en texto legible.
    - Caja      : arqueos (turnos) con su cuadre.
    - Clientes  : clientes con saldo + movimientos de cuenta corriente.

Librería: openpyxl (estándar de facto para .xlsx, ya usado por el proyecto en
`app/routers/productos.py`). Un único archivo binario, sin dependencias pesadas.
"""

import io
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.dependencies import UsuarioActual, requiere_rol
from app.models import (
    Arqueo,
    Cliente,
    MovimientoCuenta,
    Producto,
    Venta,
    VentaItem,
)

router = APIRouter(tags=["Exportar"])


def _num(valor) -> float:
    """Convierte Decimal/None a float seguro para una celda de Excel."""
    if valor is None:
        return 0.0
    try:
        return float(Decimal(valor))
    except Exception:  # noqa: BLE001
        return 0.0


def _fecha(valor) -> str:
    """Formatea datetime/date a 'AAAA-MM-DD HH:MM:SS' (o '' si es None)."""
    if valor is None:
        return ""
    try:
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:  # noqa: BLE001
        return str(valor)


def _auto_ancho(ws) -> None:
    """Ajusta el ancho de cada columna al contenido (cap para no excederse)."""
    for columna in ws.columns:
        letra = columna[0].column_letter
        largo = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
        ws.column_dimensions[letra].width = min(max(largo + 2, 12), 45)


@router.get(
    "/exportar-todo",
    summary="Exportar TODOS los datos del comercio a Excel (solo admin)",
)
def exportar_todo(
    usuario: UsuarioActual = Depends(requiere_rol("admin")),
    db: Session = Depends(get_db),
):
    """Genera un .xlsx con todas las hojas de respaldo del negocio.

    Hojas: Productos, Ventas, Caja, Clientes.
    Acceso: solo administrador. Disponible aun con licencia vencida.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        negrita = Font(bold=True)

        def _cabecera(ws, columnas):
            ws.append(columnas)
            for celda in ws[1]:
                celda.font = negrita

        # ---------------------------------------------------------------
        # HOJA 1 — PRODUCTOS (catálogo completo)
        # ---------------------------------------------------------------
        ws_prod = wb.active
        ws_prod.title = "Productos"
        _cabecera(
            ws_prod,
            [
                "nombre",
                "codigo_barras",
                "categoria",
                "marca",
                "talle",
                "color",
                "tipo_unidad",
                "precio",
                "costo",
                "porcentaje_ganancia",
                "stock",
                "stock_minimo",
                "fecha_vencimiento",
                "activo",
            ],
        )
        for p in db.query(Producto).order_by(Producto.nombre.asc()).all():
            ws_prod.append(
                [
                    p.nombre,
                    p.codigo_barras or "",
                    p.categoria or "",
                    p.marca or "",
                    p.talle or "",
                    p.color or "",
                    p.tipo_unidad or "unidad",
                    _num(p.precio),
                    _num(p.costo),
                    _num(p.porcentaje_ganancia),
                    _num(p.stock),
                    _num(p.stock_minimo),
                    p.fecha_vencimiento.strftime("%Y-%m-%d")
                    if p.fecha_vencimiento
                    else "",
                    1 if p.activo else 0,
                ]
            )
        _auto_ancho(ws_prod)

        # ---------------------------------------------------------------
        # HOJA 2 — VENTAS (cabecera + detalle legible de artículos)
        # ---------------------------------------------------------------
        ws_ventas = wb.create_sheet("Ventas")
        _cabecera(
            ws_ventas,
            [
                "folio",
                "fecha",
                "total",
                "metodo_pago",
                "cliente",
                "estado",
                "paga_con",
                "vuelto",
                "articulos",
                "anulada_en",
                "anulada_por",
                "motivo_anulacion",
            ],
        )
        productos_por_id = {p.id: p.nombre for p in db.query(Producto).all()}
        ventas = db.query(Venta).order_by(Venta.fecha.asc()).all()
        for v in ventas:
            items = db.query(VentaItem).filter(VentaItem.venta_id == v.id).all()
            detalle = "; ".join(
                f"{_num(i.cantidad)}x {productos_por_id.get(i.producto_id, '¿?')}"
                for i in items
            )
            ws_ventas.append(
                [
                    v.folio,
                    _fecha(v.fecha),
                    _num(v.total),
                    v.metodo_pago or "",
                    v.cliente or "",
                    v.estado or "activa",
                    _num(v.paga_con) if v.paga_con is not None else "",
                    _num(v.vuelto) if v.vuelto is not None else "",
                    detalle,
                    _fecha(v.anulada_en),
                    v.anulada_por or "",
                    v.motivo_anulacion or "",
                ]
            )
        _auto_ancho(ws_ventas)

        # ---------------------------------------------------------------
        # HOJA 3 — CAJA (arqueos = turnos, con su cuadre)
        # ---------------------------------------------------------------
        ws_caja = wb.create_sheet("Caja")
        _cabecera(
            ws_caja,
            [
                "fecha_apertura",
                "fecha_cierre",
                "estado",
                "monto_inicial",
                "monto_esperado",
                "monto_real",
                "diferencia",
                "observaciones_apertura",
                "observaciones_cierre",
            ],
        )
        arqueos = db.query(Arqueo).order_by(Arqueo.fecha_apertura.asc()).all()
        for a in arqueos:
            ws_caja.append(
                [
                    _fecha(a.fecha_apertura),
                    _fecha(a.fecha_cierre),
                    a.estado or "",
                    _num(a.monto_inicial),
                    _num(a.monto_esperado) if a.monto_esperado is not None else "",
                    _num(a.monto_real) if a.monto_real is not None else "",
                    _num(a.diferencia) if a.diferencia is not None else "",
                    a.observaciones_apertura or "",
                    a.observaciones_cierre or "",
                ]
            )
        _auto_ancho(ws_caja)

        # ---------------------------------------------------------------
        # HOJA 4 — CLIENTES (saldos + movimientos de cuenta corriente)
        # ---------------------------------------------------------------
        ws_cli = wb.create_sheet("Clientes")
        _cabecera(
            ws_cli,
            [
                "tipo",
                "nombre",
                "telefono",
                "saldo_deuda",
                "fecha",
                "mov_tipo",
                "mov_monto",
                "mov_concepto",
            ],
        )
        for c in db.query(Cliente).order_by(Cliente.nombre.asc()).all():
            movs = (
                db.query(MovimientoCuenta)
                .filter(MovimientoCuenta.cliente_id == c.id)
                .order_by(MovimientoCuenta.fecha.asc())
                .all()
            )
            if not movs:
                # Cliente sin movimientos: una fila con sus datos y saldo.
                ws_cli.append(
                    [
                        "cliente",
                        c.nombre,
                        c.telefono or "",
                        _num(c.saldo_deuda),
                        "",
                        "",
                        "",
                        "",
                    ]
                )
            for m in movs:
                ws_cli.append(
                    [
                        "cliente",
                        c.nombre,
                        c.telefono or "",
                        _num(c.saldo_deuda),
                        _fecha(m.fecha),
                        m.tipo or "",
                        _num(m.monto),
                        m.concepto or "",
                    ]
                )
        _auto_ancho(ws_cli)

        # ---------------------------------------------------------------
        # Serializar el libro a memoria y streamear la descarga.
        # ---------------------------------------------------------------
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        headers = {
            "Content-Disposition": "attachment; filename=respaldo_vortex_apex.xlsx",
        }
        return StreamingResponse(
            buffer,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers=headers,
        )
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error al exportar los datos: {str(exc)}",
        )
