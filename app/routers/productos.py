"""Router de operaciones CRUD para Productos y consulta de stock bajo.

Persistencia local vía SQLAlchemy + SQLite. El `db` se inyecta con
`Depends(get_db)` y las operaciones usan los modelos ORM (`app.models.Producto`).
"""

import io
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import List, Optional
from uuid import UUID

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Query,
    UploadFile,
    status,
)
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.errors import detalle_error_interno
from app.database import get_db
from app.dependencies import get_current_user, requiere_rol, UsuarioActual

from app.models import Producto, now_utc
from app.schemas.producto import ProductoCreate, ProductoUpdate, ProductoOut

router = APIRouter(prefix="/productos", tags=["Productos"])

def _sin_costos(producto: Producto, usuario: UsuarioActual) -> ProductoOut:
    """Producto listo para devolver: un cajero no recibe costo ni margen.

    Se arma una copia del esquema de salida; el objeto ORM no se toca (mutarlo
    podría guardarse sin querer en el próximo commit de la sesión).
    """
    salida = ProductoOut.model_validate(producto)
    if usuario.rol != "admin":
        salida = salida.model_copy(
            update={"costo": None, "porcentaje_ganancia": None}
        )
    return salida


# Límites de la importación desde Excel (evita agotar la memoria con un archivo
# gigante o comprimido a propósito).
_MAX_IMPORT_BYTES = 10 * 1024 * 1024
_MAX_IMPORT_FILAS = 20000

# Un texto que empieza con estos caracteres lo interpreta Excel como fórmula.
_INICIOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def _texto_seguro(valor):
    """Neutraliza fórmulas al exportar: antepone ' a un texto que parezca una."""
    if isinstance(valor, str) and valor.startswith(_INICIOS_FORMULA):
        return "'" + valor
    return valor


def _quitar_guarda_formula(valor):
    """Inversa de `_texto_seguro` al importar (el ' lo puso nuestra exportación)."""
    if (
        isinstance(valor, str)
        and len(valor) > 1
        and valor[0] == "'"
        and valor[1:].startswith(_INICIOS_FORMULA)
    ):
        return valor[1:]
    return valor

# NOTA: SQLAlchemy devuelve las columnas Numeric como `Decimal` de forma nativa;
# ya no hace falta normalizar manualmente como ocurría con Supabase.


@router.get("", response_model=List[ProductoOut], summary="Listar productos")
def listar_productos(
    categoria: str | None = None,
    activo: bool | None = None,
    estado_stock: str | None = None,  # 'bajo' | 'agotado' | 'disponible'
    orden: str = "nombre",  # 'nombre' | 'precio' | 'stock' | 'categoria'
    dir_: str = Query("asc", alias="dir"),  # asc | desc (URI: ?dir=asc)
    usuario: UsuarioActual = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Devuelve el catálogo de productos con filtros + orden (toolbar inventario).

    Filtros opcionales:
      - categoria   : coincidencia exacta con la categoría textual del producto.
      - activo      : true/false → sólo activos/inactivos.
      - estado_stock: 'bajo' (stock <= mínimo), 'agotado' (stock == 0),
                      'disponible' (resto, sin bajo ni agotado).
    Orden (whitelist para impedir SQLi):
      - orden: nombre | precio | stock | categoria
      - dir : asc | desc
    """
    try:
        query = db.query(Producto)
        if categoria:
            query = query.filter(Producto.categoria == categoria)
        if activo is not None:
            query = query.filter(Producto.activo == activo)

        if estado_stock == "bajo":
            query = query.filter(Producto.stock <= Producto.stock_minimo)
        elif estado_stock == "agotado":
            query = query.filter(Producto.stock <= 0)
        elif estado_stock == "disponible":
            query = query.filter(
                ~(Producto.stock <= Producto.stock_minimo), Producto.stock > 0
            )

        # ---------- Orden con whitelist de columnas (anti inyección) ----------
        col_orden = {
            "nombre": Producto.nombre,
            "precio": Producto.precio,
            "stock": Producto.stock,
            "categoria": Producto.categoria,
        }.get((orden or "nombre").lower(), Producto.nombre)

        col_orden = (
            col_orden.desc() if (dir_ or "asc").lower() == "desc" else col_orden.asc()
        )
        query = query.order_by(col_orden, Producto.nombre.asc())
        return [_sin_costos(p, usuario) for p in query.all()]
    except Exception as exc:  # pragma: no cover
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al listar productos", exc),
        )


@router.get(
    "/bajo-stock",
    response_model=List[ProductoOut],
    summary="Consultar alerta de bajo stock",
)
def productos_bajo_stock(
    usuario: UsuarioActual = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Devuelve los productos cuyo stock actual es menor o igual al stock mínimo."""
    try:
        productos = (
            db.query(Producto)
            .filter(Producto.activo.is_(True), Producto.stock <= Producto.stock_minimo)
            .all()
        )
        return [_sin_costos(p, usuario) for p in productos]
    except Exception as exc:  # pragma: no cover
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al consultar bajo stock", exc),
        )


@router.get(
    "/codigo/{codigo_barras}",
    response_model=ProductoOut,
    summary="Buscar producto por código de barras",
)
def buscar_por_codigo(
    codigo_barras: str,
    usuario: UsuarioActual = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Busca un producto activo por su código de barras/SKU.

    Especialmente útil en el punto de venta con lector de código de barras.
    """
    try:
        producto = (
            db.query(Producto)
            .filter(
                Producto.codigo_barras == codigo_barras,
                Producto.activo.is_(True),
            )
            .first()
        )
        if producto is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Producto no encontrado para ese código de barras",
            )
        return _sin_costos(producto, usuario)
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al buscar por código de barras", exc),
        )


@router.patch(
    "/{producto_id}/stock",
    response_model=ProductoOut,
    summary="Reponer / ajustar stock",
)
def ajustar_stock(
    producto_id: UUID,
    cantidad: Decimal = Query(..., ge=-999_999_999, le=999_999_999),
    usuario: UsuarioActual = Depends(requiere_rol("admin")),
    db: Session = Depends(get_db),
):
    """Ajusta el stock de un producto (reposición de inventario).

    Envía `cantidad`, que puede ser positiva (reponer) o negativa (baja/ajuste).
    Para productos por peso admite fracciones (Decimal); el stock nunca puede
    quedar negativo.
    """
    try:
        producto = db.query(Producto).filter(Producto.id == producto_id).first()
        if producto is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Producto no encontrado",
            )
        nuevo_stock = Decimal(producto.stock) + cantidad
        if nuevo_stock < 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    f"No puede quedar stock negativo para '{producto.nombre}' "
                    f"(stock actual: {producto.stock}, ajuste: {cantidad:+})"
                ),
            )
        producto.stock = nuevo_stock
        producto.actualizado_en = now_utc()
        db.commit()
        db.refresh(producto)
        return producto
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al ajustar stock", exc),
        )


# ========================= EXPORTAR / IMPORTAR XLSX =========================
# Columnas del Excel de inventario (orden y encabezado estable, ES).
COLUMNAS_XLSX = [
    "nombre",
    "codigo_barras",
    "categoria",
    "marca",
    "talle",
    "color",
    "tipo_papel",
    "color_bn",
    "tipo_unidad",
    "precio",
    "costo",
    "porcentaje_ganancia",
    "stock",
    "stock_minimo",
    "fecha_vencimiento",
    "activo",
]


@router.get(
    "/exportar",
    summary="Exportar el catálogo de productos a Excel (.xlsx)",
)
def exportar_productos(
    # Solo admin: la planilla incluye costo y margen de cada producto.
    usuario: UsuarioActual = Depends(requiere_rol("admin")),
    db: Session = Depends(get_db),
):
    """Genera y descarga un .xlsx con TODO el catálogo (activos e inactivos).

    Pensado para respaldo, edición masiva y carga entre instalaciones. El
    archivo resultante es re-importable tal cual (columnas idénticas).
    """
    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Productos"
        ws.append(COLUMNAS_XLSX)

        productos = db.query(Producto).order_by(Producto.nombre.asc()).all()
        for p in productos:
            ws.append(
                [
                    _texto_seguro(p.nombre),
                    _texto_seguro(p.codigo_barras or ""),
                    _texto_seguro(p.categoria or ""),
                    _texto_seguro(p.marca or ""),
                    _texto_seguro(p.talle or ""),
                    _texto_seguro(p.color or ""),
                    _texto_seguro(p.tipo_papel or ""),
                    _texto_seguro(p.color_bn or ""),
                    p.tipo_unidad or "unidad",
                    float(p.precio or 0),
                    float(p.costo or 0),
                    float(p.porcentaje_ganancia or 0),
                    float(p.stock or 0),
                    float(p.stock_minimo or 0),
                    (
                        p.fecha_vencimiento.strftime("%Y-%m-%d")
                        if p.fecha_vencimiento
                        else ""
                    ),
                    1 if p.activo else 0,
                ]
            )

        # Anchos de columna legibles.
        for i, _ in enumerate(COLUMNAS_XLSX, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        headers = {
            "Content-Disposition": "attachment; filename=inventario.xlsx",
        }
        return StreamingResponse(
            buffer,
            media_type=(
                "application/vnd.openxmlformats-officedocument." "spreadsheetml.sheet"
            ),
            headers=headers,
        )
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al exportar productos", exc),
        )


def _a_decimal(valor, default: str = "0") -> Decimal:
    """Convierte una celda de Excel a Decimal de forma tolerante."""
    if valor is None or valor == "":
        return Decimal(default)
    try:
        return Decimal(str(valor).replace(",", ".").strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f"valor numérico inválido: {valor!r}")


def _a_fecha(valor) -> Optional[datetime]:
    """Convierte una celda a datetime (acepta date, datetime o 'YYYY-MM-DD')."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor
    try:
        return datetime.strptime(str(valor).strip()[:10], "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"fecha inválida (use AAAA-MM-DD): {valor!r}")


@router.post(
    "/importar",
    summary="Importar/actualizar productos desde Excel (.xlsx) (solo admin)",
)
async def importar_productos(
    archivo: UploadFile = File(..., description="Archivo .xlsx de inventario"),
    usuario: UsuarioActual = Depends(requiere_rol("admin")),
    db: Session = Depends(get_db),
):
    """Carga masiva de catálogo desde .xlsx (upsert por código de barras).

    Estrategia (FIRME pero NUNCA destructiva):
      - Si la fila trae `codigo_barras` y YA existe un producto con ese código,
        se ACTUALIZA (upsert) — clave natural del comercio.
      - Si no existe (o no trae código), se CREA un producto nuevo.
      - El import es TOLERANTE: las filas con error se reportan y se omiten;
        las válidas se aplican. NO revierte todo por una fila mala.
      - Se aplica todo en una transacción por lote (commit al final).

    Devuelve un resumen: creados, actualizados, omitidos y detalle de errores.
    """
    nombre_archivo = (archivo.filename or "").lower()
    if not nombre_archivo.endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Formato no soportado. Subí un archivo .xlsx de Excel.",
        )

    try:
        from openpyxl import load_workbook

        # Se lee 1 byte de más para detectar el exceso sin cargar todo el archivo.
        contenido = await archivo.read(_MAX_IMPORT_BYTES + 1)
        if len(contenido) > _MAX_IMPORT_BYTES:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail=f"El archivo supera el máximo de {_MAX_IMPORT_BYTES // (1024 * 1024)} MB.",
            )
        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        ws = wb.active

        filas = []
        for fila_xl in ws.iter_rows(values_only=True):
            filas.append(fila_xl)
            if len(filas) > _MAX_IMPORT_FILAS + 1:
                raise HTTPException(
                    status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                    detail=f"El archivo tiene más de {_MAX_IMPORT_FILAS} filas.",
                )
        if not filas:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El archivo está vacío.",
            )

        # Mapear encabezados -> índice de columna (case-insensitive, tolerante).
        encabezados = [
            str(c).strip().lower() if c is not None else "" for c in filas[0]
        ]
        idx = {col: i for i, col in enumerate(encabezados) if col}

        def celda(fila, col):
            i = idx.get(col)
            valor = fila[i] if (i is not None and i < len(fila)) else None
            return _quitar_guarda_formula(valor)

        creados = 0
        actualizados = 0
        omitidos = 0
        errores: list[dict] = []

        for num, fila in enumerate(filas[1:], start=2):  # fila 1 = encabezado
            if not any(c not in (None, "") for c in fila):
                continue  # fila totalmente vacía
            try:
                nombre = (
                    (celda(fila, "nombre") or "").strip()
                    if isinstance(celda(fila, "nombre"), str)
                    else celda(fila, "nombre")
                )
                if not nombre:
                    raise ValueError("'nombre' es obligatorio")

                precio = _a_decimal(celda(fila, "precio"))
                if precio <= 0:
                    raise ValueError("'precio' debe ser mayor a cero")

                tipo = (str(celda(fila, "tipo_unidad") or "unidad")).strip().lower()
                if tipo not in Producto.TIPOS_UNIDAD:
                    tipo = "unidad"

                codigo = celda(fila, "codigo_barras")
                codigo = str(codigo).strip() if codigo not in (None, "") else None

                datos = {
                    "nombre": nombre,
                    "categoria": (celda(fila, "categoria") or None),
                    "marca": (celda(fila, "marca") or None),
                    "talle": (celda(fila, "talle") or None),
                    "color": (celda(fila, "color") or None),
                    "tipo_papel": (celda(fila, "tipo_papel") or None),
                    "color_bn": (celda(fila, "color_bn") or None),
                    "tipo_unidad": tipo,
                    "precio": precio,
                    "costo": _a_decimal(celda(fila, "costo")),
                    "porcentaje_ganancia": _a_decimal(
                        celda(fila, "porcentaje_ganancia")
                    ),
                    "stock": _a_decimal(celda(fila, "stock")),
                    "stock_minimo": _a_decimal(celda(fila, "stock_minimo")),
                    "fecha_vencimiento": _a_fecha(celda(fila, "fecha_vencimiento")),
                    "activo": str(celda(fila, "activo") or "1").strip()
                    not in ("0", "false", "False", "no", "NO", ""),
                }

                existente = (
                    db.query(Producto).filter(Producto.codigo_barras == codigo).first()
                    if codigo
                    else None
                )

                if existente is not None:
                    for k, v in datos.items():
                        setattr(existente, k, v)
                    existente.actualizado_en = now_utc()
                    actualizados += 1
                else:
                    db.add(Producto(codigo_barras=codigo, **datos))
                    creados += 1
            except ValueError as ve:
                omitidos += 1
                errores.append({"fila": num, "error": str(ve)})
            except Exception as e:  # fila problemática: reportar y continuar
                omitidos += 1
                errores.append({"fila": num, "error": str(e)})

        db.commit()
        return {
            "creados": creados,
            "actualizados": actualizados,
            "omitidos": omitidos,
            "errores": errores[:50],  # tope para no devolver respuestas gigantes
            "mensaje": (
                f"Importación finalizada: {creados} creados, "
                f"{actualizados} actualizados, {omitidos} omitidos."
            ),
        }
    except HTTPException:
        db.rollback()
        raise
    except Exception as exc:  # pragma: no cover
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al importar productos", exc),
        )


@router.get("/{producto_id}", response_model=ProductoOut, summary="Obtener un producto")
def obtener_producto(
    producto_id: UUID,
    usuario: UsuarioActual = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Devuelve el detalle de un producto específico por su ID."""
    try:
        producto = db.query(Producto).filter(Producto.id == producto_id).first()
        if producto is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Producto no encontrado",
            )
        return _sin_costos(producto, usuario)
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al obtener el producto", exc),
        )


@router.post(
    "",
    response_model=ProductoOut,
    status_code=status.HTTP_201_CREATED,
    summary="Crear un producto",
)
def crear_producto(
    producto: ProductoCreate,
    usuario: UsuarioActual = Depends(requiere_rol("admin")),
    db: Session = Depends(get_db),
):
    """Registra un nuevo producto en el catálogo."""
    try:
        nuevo = Producto(**producto.model_dump())
        db.add(nuevo)
        db.commit()
        db.refresh(nuevo)
        return nuevo
    except Exception as exc:  # pragma: no cover
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al crear el producto (posible duplicado de código de barras)", exc),
        )


@router.patch(
    "/{producto_id}", response_model=ProductoOut, summary="Editar un producto"
)
def editar_producto(
    producto_id: UUID,
    producto: ProductoUpdate,
    usuario: UsuarioActual = Depends(requiere_rol("admin")),
    db: Session = Depends(get_db),
):
    """Actualiza uno o más campos de un producto existente."""
    try:
        target = db.query(Producto).filter(Producto.id == producto_id).first()
        if target is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Producto no encontrado",
            )
        datos = {k: v for k, v in producto.model_dump().items() if v is not None}
        if not datos:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No se enviaron campos para actualizar",
            )
        for campo, valor in datos.items():
            setattr(target, campo, valor)
        target.actualizado_en = now_utc()
        db.commit()
        db.refresh(target)
        return target
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al editar el producto", exc),
        )


@router.delete(
    "/{producto_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    summary="Eliminar un producto",
)
def eliminar_producto(
    producto_id: UUID,
    usuario: UsuarioActual = Depends(requiere_rol("admin")),
    db: Session = Depends(get_db),
):
    """Elimina un producto del catálogo de forma definitiva."""
    try:
        target = db.query(Producto).filter(Producto.id == producto_id).first()
        if target is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Producto no encontrado",
            )
        db.delete(target)
        db.commit()
        return None
    except HTTPException:
        raise
    except Exception as exc:  # pragma: no cover
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=detalle_error_interno("Error al eliminar el producto", exc),
        )
